from __future__ import annotations

from collections import Counter, defaultdict
from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "fulfillment-inbound-analysis.xlsx"
PRESENTATION = ROOT / "presentation" / "international-fulfillment-reconciliation.pdf"


def business_days_to_ship(order_date, ship_date) -> int | None:
    """Replicates Excel NETWORKDAYS(order_date, ship_date) - 1 with no holidays."""
    if order_date is None or ship_date is None:
        return None

    current = order_date.date()
    end = ship_date.date()
    count = 0

    while current <= end:
        if current.weekday() < 5:
            count += 1
        current += timedelta(days=1)

    return max(0, count - 1)


def classify_sla(order_date, ship_date) -> tuple[str, int | None]:
    days = business_days_to_ship(order_date, ship_date)
    if ship_date is None:
        return "Missing Ship Date", days
    if days is not None and days > 2:
        return "Over SLA", days
    return "Within SLA", days


def load_rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in row):
            rows.append(dict(zip(headers, row)))
    return rows


def validate() -> dict:
    workbook = load_workbook(WORKBOOK, data_only=True)
    orders = load_rows(workbook["Raw Order Data"])
    inbound = load_rows(workbook["Inbound Data"])

    sla_counts = Counter()
    warehouse_counts = defaultdict(lambda: Counter())
    short_shipments = []

    for row in orders:
        status, days = classify_sla(row["Order Date"], row["Ship Date"])
        sla_counts[status] += 1

        warehouse = row["Warehouse"]
        warehouse_counts[warehouse]["total_orders"] += 1
        if row["Ship Date"] is not None:
            warehouse_counts[warehouse]["dated_orders"] += 1
            if status == "Over SLA":
                warehouse_counts[warehouse]["over_sla"] += 1

        if row["Units Shipped"] < row["Units Ordered"]:
            short_shipments.append(
                {
                    "order_id": row["Order ID"],
                    "warehouse": warehouse,
                    "units_ordered": int(row["Units Ordered"]),
                    "units_shipped": int(row["Units Shipped"]),
                    "sla_status": status,
                    "business_days_to_ship": days,
                }
            )

    dated_orders = sla_counts["Within SLA"] + sla_counts["Over SLA"]
    breach_rate = sla_counts["Over SLA"] / dated_orders

    inbound_discrepancies = [
        row
        for row in inbound
        if row["Expected Inbound Units"] != row["Received Inbound Units"]
    ]
    expected_units = sum(row["Expected Inbound Units"] for row in inbound)
    received_units = sum(row["Received Inbound Units"] for row in inbound)
    missing_units = sum(
        max(0, row["Expected Inbound Units"] - row["Received Inbound Units"])
        for row in inbound
    )

    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(str(PRESENTATION)).pages
    )

    results = {
        "total_orders": len(orders),
        "orders_with_shipping_dates": dated_orders,
        "within_sla": sla_counts["Within SLA"],
        "over_sla": sla_counts["Over SLA"],
        "missing_ship_dates": sla_counts["Missing Ship Date"],
        "sla_breach_rate": breach_rate,
        "warehouse_performance": {
            warehouse: {
                "over_sla": counts["over_sla"],
                "dated_orders": counts["dated_orders"],
                "breach_rate": counts["over_sla"] / counts["dated_orders"],
            }
            for warehouse, counts in warehouse_counts.items()
        },
        "short_shipments": short_shipments,
        "purchase_orders_reviewed": len(inbound),
        "expected_units": int(expected_units),
        "received_units": int(received_units),
        "missing_units": int(missing_units),
        "pos_with_discrepancies": len(inbound_discrepancies),
        "po_discrepancy_rate": len(inbound_discrepancies) / len(inbound),
        "inbound_discrepancies": [
            {
                "po": row["Inbound PO"],
                "expected_units": int(row["Expected Inbound Units"]),
                "received_units": int(row["Received Inbound Units"]),
                "missing_units": int(
                    row["Expected Inbound Units"] - row["Received Inbound Units"]
                ),
                "notes": row["Operational Notes"],
            }
            for row in inbound_discrepancies
        ],
        "pdf_contains_incorrect_po_de_11056": "PO-DE-11056" in pdf_text,
        "pdf_contains_correct_po_de_1056": "PO-DE-1056" in pdf_text,
    }

    assert results["total_orders"] == 60
    assert results["orders_with_shipping_dates"] == 57
    assert results["within_sla"] == 37
    assert results["over_sla"] == 20
    assert results["missing_ship_dates"] == 3
    assert round(results["sla_breach_rate"], 3) == 0.351

    expected_warehouse = {
        "Devlnd UK": (9, 19, 0.474),
        "InterFulfillment CA": (6, 19, 0.316),
        "Fidelity NL": (5, 19, 0.263),
    }
    for warehouse, (over_sla, dated_orders, rate) in expected_warehouse.items():
        observed = results["warehouse_performance"][warehouse]
        assert observed["over_sla"] == over_sla
        assert observed["dated_orders"] == dated_orders
        assert round(observed["breach_rate"], 3) == rate

    assert [row["order_id"] for row in short_shipments] == [
        "ORD-0007",
        "ORD-0019",
        "ORD-0034",
        "ORD-0048",
    ]
    assert short_shipments[-1]["warehouse"] == "InterFulfillment CA"

    assert results["purchase_orders_reviewed"] == 4
    assert results["expected_units"] == 400
    assert results["received_units"] == 392
    assert results["missing_units"] == 8
    assert results["pos_with_discrepancies"] == 2
    assert results["po_discrepancy_rate"] == 0.5
    assert [row["po"] for row in results["inbound_discrepancies"]] == [
        "PO-FI-1022",
        "PO-DE-1056",
    ]

    return results


def main() -> None:
    results = validate()
    print("Validation passed")
    print(f"Total orders: {results['total_orders']}")
    print(f"Confirmed SLA breach rate: {results['sla_breach_rate']:.1%}")
    print(f"Inbound missing units: {results['missing_units']}")
    print(
        "PDF discrepancy detected: "
        f"PO-DE-11056={results['pdf_contains_incorrect_po_de_11056']}, "
        f"PO-DE-1056={results['pdf_contains_correct_po_de_1056']}"
    )
    print(
        "Short shipment warehouse check: "
        f"ORD-0048={results['short_shipments'][-1]['warehouse']}"
    )


if __name__ == "__main__":
    main()
